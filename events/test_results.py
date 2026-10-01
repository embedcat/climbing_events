from datetime import date

import os

from django.urls import reverse
from django.test import override_settings
from openpyxl import load_workbook

from config import settings
from events import services, xl_tools
from events.models import Event, Participant
from events.tests import ClimbingEventsBaseTestCase


def accents(**marked):
    """Результаты по всем 10 трассам: r1=1 — flash на первой, r3={'top': 2, 'zone': 1} — для французской системы."""
    result = {str(i): {'top': 0, 'zone': 0} for i in range(10)}
    for key, value in marked.items():
        top, zone = (value['top'], value['zone']) if isinstance(value, dict) else (value, value)
        result[str(int(key[1:]) - 1)] = {'top': top, 'zone': zone}
    return result


class ResultsTestBase(ClimbingEventsBaseTestCase):
    def setUp(self):
        super().setUp()
        self.event = services.create_event(owner=self.superuser, title='Results Event', date=date(2026, 10, 1))
        self.event.is_published = True
        self.event.is_results_allowed = True
        self.event.is_enter_result_allowed = True
        self.event.score_type = Event.SCORE_PROPORTIONAL
        self.event.save()

    def set_event(self, **fields):
        for name, value in fields.items():
            setattr(self.event, name, value)
        self.event.save()

    def make(self, last, first='Имя', gender=Participant.GENDER_MALE, group=0, result=None, **extra):
        participant = Participant.objects.create(
            event=self.event, last_name=last, first_name=first, gender=gender, group_index=group,
            pin=extra.pop('pin', 1000 + Participant.objects.count()), **extra)
        if result is not None:
            services.enter_results(event=self.event, participant=participant, accents=result)
            participant.refresh_from_db()
        return participant

    def payload(self, **kwargs):
        return services.get_results_payload(self.event, **kwargs)

    def table(self, gender=Participant.GENDER_MALE, group=0, **kwargs):
        return next(t for t in self.payload(**kwargs)['tables'] if t['gender'] == gender and t['group_index'] == group)

    @property
    def url(self):
        return reverse('api_events-results', args=[self.event.id])


class PlacesTests(ResultsTestBase):
    def test_participants_without_results_have_no_place(self):
        strong = self.make('Андреев', result=accents(r1=1, r2=2))
        weak = self.make('Борисов', result=accents(r2=2))
        waiting = self.make('Власов')
        for p in (strong, weak, waiting):
            p.refresh_from_db()
        self.assertEqual((strong.place, weak.place, waiting.place), (1, 2, 0))

    def test_entered_result_with_zero_score_still_gets_a_place(self):
        none_climbed = self.make('Андреев', result=accents())
        waiting = self.make('Борисов')
        services.update_results(self.event)
        none_climbed.refresh_from_db()
        waiting.refresh_from_db()
        self.assertEqual(none_climbed.place, 1)
        self.assertEqual(waiting.place, 0)

    def test_equal_scores_share_a_place(self):
        a = self.make('Андреев', result=accents(r1=2))
        b = self.make('Борисов', result=accents(r1=2))
        c = self.make('Власов', result=accents())
        for p in (a, b, c):
            p.refresh_from_db()
        self.assertEqual((a.place, b.place, c.place), (1, 1, 3))

    def test_clearing_results_clears_places(self):
        p = self.make('Андреев', result=accents(r1=2))
        self.assertEqual(p.place, 1)
        services.clear_results(self.event)
        p.refresh_from_db()
        self.assertEqual((p.place, p.is_entered_result), (0, False))


class ExcelProtocolTests(ResultsTestBase):
    def test_participants_without_result_have_no_place_in_the_protocol(self):
        self.set_event(is_count_only_entered_results=False)
        self.make('Андреев', result=accents(r1=2))
        self.make('Борисов')
        xl_tools.export_result(event=self.event)
        name = services.get_list_of_protocols(self.event)[0]['name']
        path = os.path.join(settings.PROTOCOLS_PATH, str(self.event.id), name)
        try:
            sheet = load_workbook(path)['_MALE']
            self.assertEqual([sheet.cell(row=9, column=2).value, sheet.cell(row=9, column=1).value], ['Андреев Имя', 1])
            self.assertEqual([sheet.cell(row=10, column=2).value, sheet.cell(row=10, column=1).value],
                             ['Борисов Имя', None])
        finally:
            services.remove_file(f"{self.event.id}/{name}")


class ResultsPayloadTests(ResultsTestBase):
    def test_tables_per_gender_and_group(self):
        self.set_event(group_num=2, group_list='Новички, Спорт')
        payload = self.payload()
        self.assertEqual(payload['groups'], ['Новички', 'Спорт'])
        self.assertEqual([(t['gender'], t['group']) for t in payload['tables']], [
            ('MALE', 'Новички'), ('MALE', 'Спорт'), ('FEMALE', 'Новички'), ('FEMALE', 'Спорт')])

    def test_single_group_event_has_one_table_per_gender_and_no_group_names(self):
        payload = self.payload()
        self.assertEqual(payload['groups'], [])
        self.assertEqual([t['gender'] for t in payload['tables']], ['MALE', 'FEMALE'])

    def test_ranked_by_score_then_waiting_without_place(self):
        self.make('Яковлев', result=accents(r1=2, r2=2))
        self.make('Андреев', result=accents(r1=2))
        self.make('Борисов', result=accents(r1=2))
        self.make('Власов')
        self.make('Гаврилов')
        table = self.table()
        self.assertEqual([r['last_name'] for r in table['ranked']], ['Яковлев', 'Андреев', 'Борисов'])
        self.assertEqual([r['place'] for r in table['ranked']], [1, 2, 2])
        self.assertEqual([r['last_name'] for r in table['waiting']], ['Власов', 'Гаврилов'])
        self.assertEqual([r['place'] for r in table['waiting']], [None, None])
        self.assertEqual(table['waiting'][0]['results'], [])

    def test_results_and_counted_routes(self):
        self.set_event(count_routes_num=1)
        self.make('Андреев', result=accents(r1=1, r2=2))
        row = self.table()['ranked'][0]
        self.assertEqual(len(row['results']), 10)
        self.assertEqual(row['results'][0], {'top': 1, 'zone': 0})
        self.assertEqual(row['results'][1], {'top': 2, 'zone': 0})
        self.assertEqual(row['counted'].count(True), 1)
        self.assertTrue(row['counted'][0])
        self.assertEqual(self.payload()['display']['best_routes_num'], 1)

    def test_best_routes_num_only_for_systems_that_use_it(self):
        self.set_event(count_routes_num=5, score_type=Event.SCORE_FRENCH)
        self.assertEqual(self.payload()['display']['best_routes_num'], 0)

    def test_route_points_follow_the_score_formula(self):
        self.make('Андреев', result=accents(r1=1))
        self.make('Борисов', result=accents(r1=2))
        points = self.table()['route_points']
        # на первой трассе двое: 80 очков делятся пополам, flash добавляет 25%
        self.assertEqual(points[0], {'flash': 50.0, 'redpoint': 40.0})
        self.assertEqual(points[1], {'flash': 0.0, 'redpoint': 0.0})

    def test_route_points_hidden_when_not_meaningful_or_not_allowed(self):
        self.set_event(score_type=Event.SCORE_FRENCH)
        self.assertIsNone(self.table()['route_points'])
        self.set_event(score_type=Event.SCORE_NUM_ACCENTS)
        self.assertIsNone(self.table()['route_points'])
        self.set_event(score_type=Event.SCORE_PROPORTIONAL, is_view_route_score=False)
        self.assertIsNone(self.table()['route_points'])

    def test_french_results_and_score_view(self):
        self.set_event(score_type=Event.SCORE_FRENCH)
        self.make('Андреев', result=accents(r1={'top': 3, 'zone': 2}, r2={'top': 0, 'zone': 4}))
        row = self.table()['ranked'][0]
        self.assertEqual(row['results'][0], {'top': 3, 'zone': 2})
        self.assertEqual(row['results'][1], {'top': 0, 'zone': 4})
        self.assertEqual(row['score_view'], '1T3 2z6')

    def test_short_results_when_full_view_is_off(self):
        self.set_event(is_view_full_results=False)
        self.make('Андреев', result=accents(r1=2))
        row = self.table()['ranked'][0]
        self.assertEqual((row['results'], row['counted']), ([], []))
        self.assertEqual(row['place'], 1)
        self.assertFalse(self.payload()['display']['is_view_full_results'])
        self.assertIsNone(self.table()['route_points'])

    def test_participant_card_data(self):
        self.set_event(set_num=2, set_list='Утро, Вечер')
        self.make('Андреев', first='Иван', result=accents(r1=1), birth_year=1990, city='Тула', team='Скала',
                  grade=Participant.GRADE_KMS, set_index=1, email='ivan@example.com', phone_number='+79990001122')
        row = self.table()['ranked'][0]
        self.assertEqual((row['birth_year'], row['city'], row['team'], row['grade'], row['set']),
                         (1990, 'Тула', 'Скала', 'КМС', 'Вечер'))

    def test_missing_optional_fields(self):
        self.make('Андреев', result=accents(r1=1), birth_year=None, city=None, team=None)
        row = self.table()['ranked'][0]
        self.assertEqual((row['birth_year'], row['city'], row['team']), (None, '', ''))

    def test_routes_show_grade_and_color_only_when_enabled(self):
        self.assertEqual(self.payload()['routes'][0], {'number': 1, 'grade': None, 'color': None})
        self.set_event(is_view_route_grade=True, is_view_route_color=True)
        route = self.payload()['routes'][0]
        self.assertEqual((route['grade'], route['color']), ('5', '#FF0000'))

    def test_event_block(self):
        data = self.payload(can_edit=True)['event']
        self.assertEqual((data['id'], data['title'], data['score_type'], data['routes_num'], data['can_edit']),
                         (self.event.id, 'Results Event', 'PROP', 10, True))

    def test_live_while_entry_is_open_and_event_not_expired(self):
        self.assertTrue(self.payload()['event']['is_live'])
        self.set_event(is_enter_result_allowed=False)
        self.assertFalse(self.payload()['event']['is_live'])
        self.set_event(is_enter_result_allowed=True, is_expired=True)
        self.assertFalse(self.payload()['event']['is_live'])
        self.assertTrue(self.payload()['event']['is_expired'])


class ResultsApiTests(ResultsTestBase):
    def test_public_results_contain_no_private_data(self):
        self.make('Андреев', result=accents(r1=1), pin=4321, email='ivan@example.com', phone_number='+79990001122')
        self.make('Борисов', pin=5432, email='boris@example.com')
        body = self.client.get(self.url).content.decode()
        for private in ('ivan@example.com', 'boris@example.com', '+79990001122', '4321', '5432', 'owner', 'admin'):
            self.assertNotIn(private, body)

    def test_anonymous_gets_published_results(self):
        self.make('Андреев', result=accents(r1=1))
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['tables'][0]['ranked'][0]['last_name'], 'Андреев')
        self.assertFalse(data['event']['can_edit'])

    def test_unpublished_event_is_hidden(self):
        self.set_event(is_published=False)
        self.assertEqual(self.client.get(self.url).status_code, 404)

    def test_results_closed(self):
        self.set_event(is_results_allowed=False)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()['code'], 'results_closed')

    def test_can_edit_only_for_owner_and_superuser(self):
        self.client.force_login(self.user)
        self.assertFalse(self.client.get(self.url).json()['event']['can_edit'])
        self.client.force_login(self.superuser)
        self.assertTrue(self.client.get(self.url).json()['event']['can_edit'])
        owner = services.create_event(owner=self.user, title='Mine', date=date(2026, 10, 2))
        owner.is_published = owner.is_results_allowed = True
        owner.save()
        self.client.force_login(self.user)
        data = self.client.get(reverse('api_events-results', args=[owner.id])).json()
        self.assertTrue(data['event']['can_edit'])

    def test_etag_lets_the_client_skip_unchanged_results(self):
        self.make('Андреев', result=accents(r1=1))
        first = self.client.get(self.url)
        etag = first['ETag']
        self.assertEqual(first['Cache-Control'], 'private, no-cache')

        again = self.client.get(self.url, HTTP_IF_NONE_MATCH=etag)
        self.assertEqual(again.status_code, 304)
        self.assertEqual(again.content, b'')
        self.assertEqual(again['ETag'], etag)

    def test_etag_changes_when_results_change(self):
        p = self.make('Андреев', result=accents(r1=1))
        etag = self.client.get(self.url)['ETag']
        services.enter_results(event=self.event, participant=p, accents=accents(r1=2))
        changed = self.client.get(self.url, HTTP_IF_NONE_MATCH=etag)
        self.assertEqual(changed.status_code, 200)
        self.assertNotEqual(changed['ETag'], etag)

    def test_etag_is_per_viewer_permissions(self):
        etag = self.client.get(self.url)['ETag']
        self.client.force_login(self.superuser)
        response = self.client.get(self.url, HTTP_IF_NONE_MATCH=etag)
        self.assertEqual(response.status_code, 200)

    def test_new_result_shows_up_and_moves_places(self):
        self.make('Андреев', result=accents(r1=2))
        late = self.make('Борисов')
        self.assertEqual([r['last_name'] for r in self.table()['waiting']], ['Борисов'])
        services.enter_results(event=self.event, participant=late, accents=accents(r1=2, r2=2))
        table = self.table()
        self.assertEqual([(r['last_name'], r['place']) for r in table['ranked']], [('Борисов', 1), ('Андреев', 2)])
        self.assertEqual(table['waiting'], [])

    def test_few_queries_regardless_of_participants(self):
        for i in range(20):
            self.make(f'Фамилия{i:02d}', result=accents(r1=1 + i % 2))
        with self.assertNumQueries(3):
            self.client.get(self.url)


@override_settings(VITE_DEV_SERVER='http://localhost:5173')
class ResultsPageTests(ResultsTestBase):
    def test_page_mounts_vue_app(self):
        # результаты — вкладка страницы события: тот же HTML, экран выбирает Vue по адресу
        response = self.client.get(reverse('results', args=[self.event.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'id="event-app" data-event-id="{self.event.id}"')
        self.assertContains(response, 'src/entries/event.ts')

    def test_unpublished_event_shows_banner(self):
        self.set_event(is_published=False)
        response = self.client.get(reverse('results', args=[self.event.id]))
        self.assertNotContains(response, 'id="event-app"')
        self.assertContains(response, 'Событие не опубликовано')

    def test_owner_sees_page_of_unpublished_event(self):
        self.set_event(is_published=False)
        self.client.force_login(self.superuser)
        response = self.client.get(reverse('results', args=[self.event.id]))
        self.assertContains(response, 'id="event-app"')
